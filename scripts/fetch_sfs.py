# -*- coding: utf-8 -*-
"""
総務省の決算カードから標準財政規模を取り込み、data-*.json に sfs キーを追加する。

使い方（GitHub Actions の workflow_dispatch から実行）:
    python scripts/fetch_sfs.py         -> 今年度分。sfs キーを更新（他の指標と同じくスライド式で履歴を保持）
    python scripts/fetch_sfs.py r1      -> 1年前（過去分）の一括バックフィル。sfs_r1 キーに直接書き込む
                                            （スライドはしない。過去分をまとめて登録する時だけ使う特別モード）

過去分を追加するときは、config.json の sfs に base_muni_r1 / pref_card_r1 を
追加しておくこと（該当年度の決算カードの番号）。

背景:
  財政調整基金の適正水準は、一般に「標準財政規模の10〜20%」という目安で語られる。
  しかし本アプリは歳出を分母にしていたため、世間の目安と直接比較できなかった。
  標準財政規模を取り込めば、同じ土俵で判断できるようになる。

データ源:
  市区町村 … 都道府県ごとの市町村決算カード（48ファイル。北海道のみ2分割）
              1シート＝1自治体。標準財政規模は 54行97列（千円単位）に固定。
              特別区23区も東京都のファイルに含まれる。
  都道府県 … 都道府県決算カード（1ファイル）

URLについて:
  決算カードのURLは毎年変わる通し番号。PDFとExcelが交互に並ぶため、
  1団体あたり2ずつ増える。config.json の sfs.base_muni に東京都のExcel番号を
  書いておけば、残り47件は計算で求まる。北海道1のみ例外（+1）。
  番号がずれた年は、404が出るので実行時に検出できる。

  ⚠️ 2026年9月に判明：この「+2ずつの計算」が成立するのは年によって変わる。
  令和6・5年度は成立を確認済みだが、令和4年度以前は確認したところ番号が
  飛んでいる箇所があった（訂正・差し替えの影響と思われる）。過去分を追加する
  ときは、必ず対象年度で東京都のExcelの番号が周辺県と綺麗に+2ずつ並んでいるか
  総務省の決算カード一覧ページで目視確認してから base_muni_rN に設定すること。
"""

import io
import json
import os
import re
import sys
import time
import urllib.request

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

SUFFIX = sys.argv[1] if len(sys.argv) > 1 else ""
if SUFFIX and not re.fullmatch(r"r[0-9]+", SUFFIX):
    print("NG: 引数は空欄か r1, r2... の形式で指定してください")
    sys.exit(1)
FIELD = "sfs_" + SUFFIX if SUFFIX else "sfs"
CFG_BASE_KEY = "base_muni_" + SUFFIX if SUFFIX else "base_muni"
CFG_PREF_KEY = "pref_card_" + SUFFIX if SUFFIX else "pref_card"

# 決算カードの都道府県の並び（総務省の一覧ページ順）
PREF_ORDER = [
    "北海道", "北海道", "青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県",
    "茨城県", "栃木県", "群馬県", "埼玉県", "千葉県", "東京都", "神奈川県", "新潟県",
    "富山県", "石川県", "福井県", "山梨県", "長野県", "岐阜県", "静岡県", "愛知県",
    "三重県", "滋賀県", "京都府", "大阪府", "兵庫県", "奈良県", "和歌山県", "鳥取県",
    "島根県", "岡山県", "広島県", "山口県", "徳島県", "香川県", "愛媛県", "高知県",
    "福岡県", "佐賀県", "長崎県", "熊本県", "大分県", "宮崎県", "鹿児島県", "沖縄県",
]
TOKYO_INDEX = 13  # PREF_ORDER 上の東京都の位置

# 決算カードのセル位置（市町村と都道府県で異なる）
MUNI_POS = (54, 84, 97)   # 行, ラベル列, 値列
PREF_POS = (37, 81, 95)

DATA_FILES = [
    "data-hokkaido-tohoku.json", "data-kanto.json", "data-chubu.json",
    "data-kinki.json", "data-chugoku-shikoku.json", "data-kyushu.json",
]


def norm(name):
    """自治体名の表記ゆれを吸収する。

    決算カードのシート名は「六ケ所村」のように大きい「ケ」を使う一方、
    アプリ側は「六ヶ所村」と小書きの「ヶ」を使っている。
    同じ扱いにしないと突合できない（該当7件）。
    """
    return (name.replace("ヶ", "ケ").replace("ヵ", "カ")
                .replace("\u3000", "").strip())


def slide_and_set(entry, prefix, new_value, start_idx=1, max_total=8):
    """履歴フィールドをスライドし、新しい主値をセットする。

    update_fiscal_data.py の同名関数と同じロジック（8データポイント＝履歴7+最新1
    に達するまでは追加、以降は最古を破棄してローリング）。sfs は他の指標と別スクリプト
    で取得しているため、ここにも同じ処理を持たせて挙動を揃えている。

    ※ new_value が現在の主値と完全に同じ場合は「新年度データではない」と判断し、
      スライドせずスキップする（同じ総務省データで誤って再実行した際の重複防止）。
    """
    old_main = entry.get(prefix)
    if new_value is not None and old_main is not None and new_value == old_main:
        return
    n = 0
    while f"{prefix}_r{start_idx + n}" in entry:
        n += 1
    if n < max_total - 1:
        if old_main is not None:
            entry[f"{prefix}_r{start_idx + n}"] = old_main
    else:
        for i in range(start_idx, start_idx + n - 1):
            entry[f"{prefix}_r{i}"] = entry.get(f"{prefix}_r{i + 1}")
        if old_main is not None:
            entry[f"{prefix}_r{start_idx + n - 1}"] = old_main
    if new_value is not None:
        entry[prefix] = new_value


def fail(msg):
    print("NG: " + msg)
    sys.exit(1)


def build_urls(base_muni):
    """東京都の番号から48ファイルのURLを計算する。"""
    urls = []
    for i in range(len(PREF_ORDER)):
        num = base_muni + (i - TOKYO_INDEX) * 2
        if i == 0:
            num += 1  # 北海道1のみ例外（北海道2と連番のため）
        urls.append((PREF_ORDER[i],
                     "https://www.soumu.go.jp/main_content/%09d.xlsx" % num))
    return urls


def download(url, tries=3):
    for t in range(tries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=90) as res:
                return res.read()
        except Exception as e:
            if t == tries - 1:
                raise
            print("    再試行 %d/%d (%s)" % (t + 2, tries, e))
            time.sleep(3)


def read_card(data, pref, pos):
    """決算カードのExcelから {(都道府県, 自治体名): 標準財政規模(億円)} を作る。

    同名の自治体が全国に複数あるため、名前だけをキーにすると上書きされる。
    （例：金山町は山形県と福島県にある）
    そのため都道府県とのペアで持つ。
    """
    row, label_col, value_col = pos
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out = {}
    skipped = 0
    for sh in wb.sheetnames:
        if sh == "目次":
            continue
        ws = wb[sh]
        if ws.cell(row, label_col).value != "標準財政規模":
            skipped += 1
            continue
        v = ws.cell(row, value_col).value
        if not isinstance(v, (int, float)):
            skipped += 1
            continue
        name = re.sub(r"^[0-9]+", "", sh).strip()
        key_pref = pref if pref else name  # 都道府県カードは自分自身が都道府県
        out[(key_pref, norm(name))] = round(v / 100000.0, 1)  # 千円 → 億円
    wb.close()
    return out, skipped


def main():
    cfg_path = os.path.join(HERE, "config.json")
    if not os.path.exists(cfg_path):
        fail("config.json が見つかりません")
    cfg = json.load(open(cfg_path, encoding="utf-8"))

    sfs_cfg = cfg.get("sfs")
    if not sfs_cfg or not sfs_cfg.get(CFG_BASE_KEY):
        fail("config.json の sfs に %s（対象年度の市町村決算カードExcelの番号）"
             "を追加してください。例: \"%s\": 1063971" % (CFG_BASE_KEY, CFG_BASE_KEY))
    base = int(sfs_cfg[CFG_BASE_KEY])
    pref_url = sfs_cfg.get(CFG_PREF_KEY)

    # --- data-*.json を先に読み込んでおく（書き込み用、かつ都道府県ごとの答え合わせ用） ---
    dbs = {}
    expected_by_pref = {}
    expected_prefs = set()
    for fn in DATA_FILES:
        path = os.path.join(ROOT, fn)
        if not os.path.exists(path):
            fail(fn + " が見つかりません")
        db = json.load(open(path, encoding="utf-8"))
        dbs[fn] = db
        for k, v in db.items():
            p = v.get("p")
            if not p:
                continue
            expected_prefs.add(p)
            if p == k:
                continue  # 都道府県自身の行は市区町村名一覧に含めない
            base_name = norm(re.sub(r"（[^）]*）$", "", k))
            expected_by_pref.setdefault(p, set()).add(base_name)

    print("■ 市町村決算カードを取得（48ファイル） [対象キー: %s]" % FIELD)
    all_sfs = {}
    for i, (pref, url) in enumerate(build_urls(base), 1):
        print("  [%2d/48] %-6s %s" % (i, pref, url.split("/")[-1]))
        try:
            data = download(url)
        except Exception as e:
            fail("%s の取得に失敗しました: %s\n"
                 "URLの番号がずれている可能性があります。総務省の決算カードページで\n"
                 "東京都のExcelの番号を確認し、config.json の sfs.%s を"
                 "直してください。" % (pref, e, CFG_BASE_KEY))
        got, skipped = read_card(data, pref, MUNI_POS)

        # --- 答え合わせ：中身が本当にこの都道府県のものか ---
        fetched_names = set(name for (p2, name) in got.keys() if p2 == pref)
        if fetched_names:
            expected = expected_by_pref.get(pref, set())
            overlap = fetched_names & expected
            rate = len(overlap) / len(fetched_names)
            if rate < 0.5:
                sample = ", ".join(list(fetched_names)[:5])
                fail(
                    "%s のはずのファイルの中身が、アプリが知っている%sの自治体名と"
                    "ほとんど一致しません（一致率 %.0f%%）。\n"
                    "番号がズレて別の都道府県のファイルを取得している可能性が高いです。\n"
                    "実際にファイルから読めた名前の例: %s\n"
                    "総務省の決算カード一覧ページで、%s前後の番号を目視確認してから"
                    "config.json の sfs.%s をやり直してください。"
                    % (pref, pref, rate * 100, sample, pref, CFG_BASE_KEY)
                )

        all_sfs.update(got)
        print("      %d件 取得（対象外シート %d）" % (len(got), skipped))

    if pref_url:
        print("■ 都道府県決算カードを取得")
        try:
            data = download(pref_url)
            got, skipped = read_card(data, "", PREF_POS)
            fetched_pref_names = set(name for (_, name) in got.keys())
            overlap = fetched_pref_names & expected_prefs
            rate = len(overlap) / len(fetched_pref_names) if fetched_pref_names else 0
            if rate < 0.5:
                fail(
                    "都道府県決算カードの中身が、アプリが知っている都道府県名と"
                    "ほとんど一致しません（一致率 %.0f%%）。pref_card のURLがずれている"
                    "可能性があります。" % (rate * 100)
                )
            all_sfs.update(got)
            print("  %d件 取得（対象外シート %d）" % (len(got), skipped))
        except Exception as e:
            print("  取得できませんでした（%s）。都道府県分はスキップします。" % e)
    else:
        print("■ 都道府県決算カードのURLが未設定のためスキップ")

    print("")
    print("合計 %d件 の標準財政規模を取得しました" % len(all_sfs))
    if len(all_sfs) < 1500:
        fail("取得件数が少なすぎます（%d件）。ファイル構造が変わった可能性があります。"
             % len(all_sfs))

    print("■ data-*.json に反映")
    total = matched = 0
    missing = []
    for fn in DATA_FILES:
        db = dbs[fn]
        n = 0
        for k, v in db.items():
            total += 1
            # アプリのキーは「金山町（山形）」のように括弧付きの場合がある
            base_name = norm(re.sub(r"（[^）]*）$", "", k))
            look = (v.get("p"), base_name)
            if look in all_sfs:
                if SUFFIX:
                    # 過去分の一括登録（バックフィル）: 直接そのキーに書き込む
                    v[FIELD] = all_sfs[look]
                else:
                    # 今年度分の通常更新: 他の指標と同じくスライドしてからセット
                    slide_and_set(v, "sfs", all_sfs[look])
                n += 1
                matched += 1
            else:
                missing.append(k)
        path = os.path.join(ROOT, fn)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(db, f, ensure_ascii=False, separators=(",", ":"))
        print("  %-30s %d件に付与" % (fn, n))

    print("")
    print("突合結果: %d / %d 件 (%.1f%%)" % (matched, total, matched / total * 100))
    if missing:
        print("未取得（先頭20件）: %s" % ", ".join(missing[:20]))
    if matched / total < 0.95:
        fail("突合率が95%を下回りました。名称の対応を確認してください。")

    print("")
    print("完了しました。")


if __name__ == "__main__":
    main()
