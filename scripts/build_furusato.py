# -*- coding: utf-8 -*-
"""
ふるさと納税データ（受入額・住民税控除額）を総務省Excelから作り、data-*.json に書き込む。
手順の説明は scripts/README.md の「【8月】ふるさと納税の更新」を参照。

■ 毎年8月の更新（ふつうはこちら）
    python scripts/build_furusato.py <Excelフォルダ>

    Excelフォルダに入れるのは、総務省「ふるさと納税に関する現況調査結果」の次の2つだけ:
      ・各自治体のふるさと納税受入額及び受入件数（平成20年度～令和X年度）
      ・各自治体の令和Y年度課税における住民税控除額等（最新の1年分）
    ファイル名は何でもよい（中身の見出しから自動で見分ける）。

    ・受入額 fuH は、受入額Excelの最新8年分で作り直す
    ・控除額 fkH は、今のデータの過去分を1年ずらして、最新年度を末尾に足す
      （同じ年度でもう一度実行した場合は、末尾を置き換えるだけ）
    ・preview/js/ui.js のグラフの年度ラベル（FURU_YEARS）も自動で書き換える
    ・scripts/furusato_config.json に「控除額の最新課税年度」を記録する

■ 全部作り直す（データが壊れたときなど）
    python scripts/build_furusato.py <Excelフォルダ> --full

    Excelフォルダに、受入額Excel と 控除額Excel 8年分 を入れておく。

■ 書き込み先
    省略すると preview/ に書き込む（本番へは「プレビューのデータを本番に反映」
    「プレビューを本番に反映」の2つのワークフローで反映する）。
    別の場所に書く場合は --target <フォルダ> を付ける。

書き込むキー（各市区町村。都道府県の行は触らない）:
    fu  : 受入額 最新年度（万円）
    fk  : 住民税控除額 最新課税年度（万円）
    fuH : 受入額 8年分（古い順）
    fkH : 住民税控除額 8年分（古い順。課税年度は受入額の翌年度）
    fuNote : 注記（scripts/furusato_notes.json から。指定取消などの事情）

【2026年9月の修正】
  旧版は Excel の市町村名だけでキーを探していたため、同名自治体（伊達市・朝日町・金山町・太子町）で
  「伊達市（北海道）」側が全年度0になっていた。今は必ず (都道府県, 市町村名) の組で突き合わせる。
  また、Excelの年度ごとの表記ゆれ（「青森」「亀田郡七飯町」、福岡県の旧名「那珂川町」）も吸収する。
"""
import glob
import json
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

DATA_FILES = [
    "data-hokkaido-tohoku.json", "data-kanto.json", "data-chubu.json",
    "data-kinki.json", "data-chugoku-shikoku.json", "data-kyushu.json",
]
YEARS = 8  # グラフに出す年数
CONFIG_FILE = os.path.join(HERE, "furusato_config.json")
NOTES_FILE = os.path.join(HERE, "furusato_notes.json")

# Excel側の表記ゆれ
PREF_FIX = {"青森": "青森県", "沖縄": "沖縄県"}
# (都道府県, Excelの名前) -> 現在の名前（市制施行などで名前が変わったのに旧名のまま載っている年がある）
RENAMED = {("福岡県", "那珂川町"): "那珂川市"}
SKIP_WORDS = ("合計", "集計")

ZEN2HAN = str.maketrans("０１２３４５６７８９", "0123456789")


def fail(msg):
    print("NG: " + msg)
    sys.exit(1)


# ---------- 年度の扱い（令和元年度 = 2019年度 のように西暦に直して比べる） ----------

def wareki_to_year(era, num):
    num = str(num).translate(ZEN2HAN)
    n = 1 if num == "元" else int(num)
    return (2018 + n) if era == "令和" else (1988 + n)


def year_label(y):
    """2018 → "H30"、2019 → "R1" （グラフの年度ラベル）"""
    return "H%d" % (y - 1988) if y <= 2018 else "R%d" % (y - 2018)


def parse_year(text):
    """「令和８年度」「平成30年度」「令和元（平成31）年度」から西暦年度を返す"""
    m = re.search(r"(令和|平成)(元|[0-9０-９]+)", str(text))
    return wareki_to_year(m.group(1), m.group(2)) if m else None


# ---------- 自治体名の突き合わせ ----------

def norm_pref(s):
    s = str(s).replace("　", "").replace(" ", "").strip()
    return PREF_FIX.get(s, s)


def norm_muni(s):
    s = str(s).replace("　", "").replace(" ", "").strip()
    # 「亀田郡七飯町」のような郡名つき表記 → 「七飯町」（上郡町・郡山市などは対象外）
    s = re.sub(r"^[^郡]{1,4}郡(?=.+[町村]$)", "", s)
    return s.replace("ヶ", "ケ").replace("ヵ", "カ")


def load_db(target):
    """{(都道府県, 素の市町村名): DBキー} を作る。DBキーは同名自治体なら「伊達市（北海道）」の形。"""
    dbs, lookup = {}, {}
    for fn in DATA_FILES:
        path = os.path.join(target, fn)
        if not os.path.exists(path):
            fail(path + " が見つかりません")
        db = json.load(open(path, encoding="utf-8"))
        dbs[fn] = db
        for k, v in db.items():
            if k == v.get("p"):
                continue  # 都道府県
            key = (norm_pref(v["p"]), norm_muni(re.sub(r"（[^）]*）$", "", k)))
            if key in lookup:
                fail("DB内で (都道府県, 名前) が重複: %s / %s" % (lookup[key], k))
            lookup[key] = k
    return dbs, lookup


def resolve(lookup, pref, muni):
    p, m = norm_pref(pref), norm_muni(muni)
    m = RENAMED.get((p, m), m)
    return lookup.get((p, m))


# ---------- Excelの読み取り ----------

def classify(path):
    """Excelの中身から種類を見分ける → ("ukeire", None) / ("kojo", 課税年度) / (None, None)"""
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return None, None
    ws = wb.worksheets[0]
    head = [v for r in ws.iter_rows(max_row=5, values_only=True) for v in r if v]
    text = " ".join(str(v) for v in head)
    if "課税" in text and "寄附金税額控除" in text:
        m = re.search(r"(令和|平成)(元|[0-9０-９]+)[^年]{0,8}年度課税", text)
        return "kojo", (wareki_to_year(m.group(1), m.group(2)) if m else None)
    if "団体名" in text and "平成20年度" in text:
        return "ukeire", None
    return None, None


def read_ukeire(path, lookup):
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    # 見出し行（「平成20年度」「令和元年度」…が並ぶ行）から、年度ごとの「金額」列を探す
    hdr = next(r for r in rows[:6] if any(v and "平成20年度" in str(v) for v in r))
    year_cols = [(parse_year(v), j) for j, v in enumerate(hdr) if v and "年度" in str(v)]
    year_cols = sorted(year_cols)[-YEARS:]
    years = [y for y, _ in year_cols]
    if years != list(range(years[0], years[0] + YEARS)):
        fail("受入額Excelの年度の並びが想定と違います: %s" % years)
    out, unmatched = {}, []
    for r in rows:
        pref, muni = r[0], r[1]
        if not pref or not muni or any(w in str(pref) for w in SKIP_WORDS):
            continue
        if not any(isinstance(r[j], (int, float)) for _, j in year_cols):
            continue  # 見出し行など
        k = resolve(lookup, pref, muni)
        if not k:
            unmatched.append("%s %s" % (pref, muni))
            continue
        out[k] = [round((r[j] or 0) / 10, 1) for _, j in year_cols]  # 千円→万円
    return years, out, unmatched


def read_kojo(path, lookup):
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    # 「ふるさと納税に係る寄附金税額控除」の控除額（※推計値含む）列（市町村民税・道府県民税の2列）
    cols = None
    for r in rows[:25]:
        c = [j for j, v in enumerate(r) if v and "推計" in str(v)]
        if len(c) == 2:
            cols = c
            break
    if not cols:
        fail(os.path.basename(path) + ": 控除額（※推計値含む）の列が見つかりません")
    out, unmatched = {}, []
    for r in rows:
        # 令和6年度課税以降は先頭に団体コード列がある
        off = 1 if (r and r[0] is not None and re.fullmatch(r"\d{5,6}", str(r[0]).strip())) else 0
        if len(r) <= max(cols) or not r[off] or not r[off + 1]:
            continue
        if not isinstance(r[cols[0]], (int, float)):
            continue  # 見出し行など
        pref, muni = r[off], r[off + 1]
        if any(w in str(pref) for w in SKIP_WORDS):
            continue
        k = resolve(lookup, pref, muni)
        if not k:
            unmatched.append("%s %s" % (pref, muni))
            continue
        out[k] = round(((r[cols[0]] or 0) + (r[cols[1]] or 0)) / 1e4, 1)  # 円→万円
    return out, unmatched


# ---------- 注記だけ反映（2026-10-01追加） ----------

def notes_only(target):
    """furusato_notes.json の注記（fuNote）だけを data-*.json に入れ直す。受入額・控除額には触らない。
    ワークフローでURL欄を空欄のまま実行したときに使う"""
    dbs, lookup = load_db(target)
    notes = json.load(open(NOTES_FILE, encoding="utf-8")) if os.path.exists(NOTES_FILE) else {}
    names = [n for n in notes if not n.startswith("_")]
    for name in names:
        if not any(name in db for db in dbs.values()):
            fail("furusato_notes.json の自治体名がDBにありません: " + name)
        for x in notes[name]:
            if not x.get("t") or not str(x.get("u", "")).startswith("https://"):
                fail("%s の注記に本文（t）か出典URL（u）がありません" % name)
    changed = []
    for fn, db in dbs.items():
        for k, v in db.items():
            if k == v.get("p"):
                continue
            before = v.get("fuNote")
            if k in notes:
                v["fuNote"] = notes[k]
            else:
                v.pop("fuNote", None)
            if v.get("fuNote") != before:
                changed.append(k)
        with open(os.path.join(target, fn), "w", encoding="utf-8") as f:
            json.dump(db, f, ensure_ascii=False, separators=(",", ":"))
    print("OK: 注記だけ反映しました。注記のある自治体 %d件、変わった自治体 %d件 %s" % (len(names), len(changed), changed))


# ---------- メイン ----------

def main():
    args = sys.argv[1:]
    full = "--full" in args
    target = os.path.join(ROOT, "preview")
    if "--target" in args:
        target = args[args.index("--target") + 1]
    pos = [a for i, a in enumerate(args) if not a.startswith("--") and (i == 0 or args[i - 1] != "--target")]
    if "--notes-only" in args:
        return notes_only(target)
    if not pos:
        fail("使い方: python scripts/build_furusato.py <Excelフォルダ> [--full] [--target <フォルダ>]"
             "／注記だけ: python scripts/build_furusato.py --notes-only")
    src = pos[0]

    config = json.load(open(CONFIG_FILE, encoding="utf-8")) if os.path.exists(CONFIG_FILE) else {}
    dbs, lookup = load_db(target)
    print("■ DB: 市区町村 %d件（書き込み先: %s）" % (len(lookup), target))

    # Excelを見分ける
    uke_path, kojo_paths = None, {}
    for p in sorted(glob.glob(os.path.join(src, "*.xlsx"))):
        kind, y = classify(p)
        if kind == "ukeire":
            uke_path = p
        elif kind == "kojo" and y:
            kojo_paths[y] = p
        print("  %s → %s" % (os.path.basename(p), {"ukeire": "受入額", "kojo": "控除額 %s課税" % (year_label(y) if y else "?")}.get(kind, "対象外")))
    if not uke_path:
        fail("受入額のExcel（平成20年度～の推移）が見つかりません")
    if not kojo_paths:
        fail("控除額のExcel（令和X年度課税における住民税控除額等）が見つかりません")

    # 受入額
    uke_years, uke, un = read_ukeire(uke_path, lookup)
    print("■ 受入額 %s〜%s年度: %d件一致 / 不一致 %d件 %s" % (year_label(uke_years[0]), year_label(uke_years[-1]), len(uke), len(un), un[:10]))
    if un:
        fail("受入額Excelに一致しない自治体があります（上の一覧を確認）")
    kojo_last = uke_years[-1] + 1  # 控除額は寄附の翌年度に課税される
    if max(kojo_paths) != kojo_last:
        fail("控除額Excelの最新年度（%s課税）が、受入額の最新年度（%s年度）の翌年度（%s課税）と一致しません"
             % (year_label(max(kojo_paths)), year_label(uke_years[-1]), year_label(kojo_last)))

    # 控除額
    kojo_years = list(range(kojo_last - YEARS + 1, kojo_last + 1))
    kojo = {}
    if full:
        lack = [year_label(y) for y in kojo_years if y not in kojo_paths]
        if lack:
            fail("--full には控除額Excelが8年分必要です。足りない課税年度: %s" % lack)
        for i, y in enumerate(kojo_years):
            k, un = read_kojo(kojo_paths[y], lookup)
            print("■ 控除額 %s課税: %d件一致 / 不一致 %d件 %s" % (year_label(y), len(k), len(un), un[:10]))
            if un:
                fail("控除額Excelに一致しない自治体があります（上の一覧を確認）")
            for key, v in k.items():
                kojo.setdefault(key, [None] * YEARS)[i] = v
    else:
        prev_last = config.get("kojo_latest_year")
        if prev_last is None:
            fail("scripts/furusato_config.json がありません。初回は --full で実行してください")
        shift = kojo_last - prev_last
        if shift not in (0, 1):
            fail("今のデータの控除額は%s課税までです。%s課税を足すには、間の年度が抜けています（--full で作り直してください）"
                 % (year_label(prev_last), year_label(kojo_last)))
        newest, un = read_kojo(kojo_paths[kojo_last], lookup)
        print("■ 控除額 %s課税: %d件一致 / 不一致 %d件 %s" % (year_label(kojo_last), len(newest), len(un), un[:10]))
        if un:
            fail("控除額Excelに一致しない自治体があります（上の一覧を確認）")
        print("  過去分は今のデータを%s" % ("1年ずらして使います" if shift else "そのまま使い、最新年度だけ置き換えます"))
        for db in dbs.values():
            for k, v in db.items():
                if k == v.get("p") or k not in newest:
                    continue
                old = v.get("fkH") or []
                if len(old) != YEARS:
                    fail("%s の fkH が%d年分ではありません（--full で作り直してください）" % (k, YEARS))
                kojo[k] = (old[1:] if shift else old[:-1]) + [newest[k]]

    notes = json.load(open(NOTES_FILE, encoding="utf-8")) if os.path.exists(NOTES_FILE) else {}
    for name in notes:
        if not name.startswith("_") and not any(name in db for db in dbs.values()):
            fail("furusato_notes.json の自治体名がDBにありません: " + name)

    # 書き込み
    missing, changed = [], 0
    for db in dbs.values():
        for k, v in db.items():
            if k == v.get("p"):
                continue
            u, ko = uke.get(k), kojo.get(k)
            if u is None or ko is None or None in ko:
                missing.append(k)
                continue
            before = json.dumps([v.get(x) for x in ("fu", "fk", "fuH", "fkH", "fuNote")], ensure_ascii=False)
            v["fu"], v["fk"], v["fuH"], v["fkH"] = u[-1], ko[-1], u, ko
            if k in notes:
                v["fuNote"] = notes[k]
            else:
                v.pop("fuNote", None)
            if json.dumps([v.get(x) for x in ("fu", "fk", "fuH", "fkH", "fuNote")], ensure_ascii=False) != before:
                changed += 1
    if missing:
        fail("データが揃わない自治体: %s" % missing[:20])

    for fn, db in dbs.items():
        with open(os.path.join(target, fn), "w", encoding="utf-8") as f:
            json.dump(db, f, ensure_ascii=False, separators=(",", ":"))

    # ui.js のグラフ年度ラベルを書き換える
    labels = [year_label(y) for y in uke_years]
    ui_path = os.path.join(target, "js", "ui.js")
    if os.path.exists(ui_path):
        s = open(ui_path, encoding="utf-8").read()
        new_line = "var FURU_YEARS = %s;" % json.dumps(labels, separators=(",", ":"))
        s2, n = re.subn(r"var FURU_YEARS = \[[^\]]*\];", new_line, s)
        if n != 1:
            fail("ui.js の FURU_YEARS が見つかりません")
        if s2 != s:
            open(ui_path, "w", encoding="utf-8").write(s2)
            print("■ ui.js のグラフ年度ラベルを %s〜%s に更新しました" % (labels[0], labels[-1]))

    config["kojo_latest_year"] = kojo_last
    config["_説明"] = "build_furusato.py が自動で書き換える。kojo_latest_year = data-*.json の控除額(fkH)の最新課税年度（西暦年度。2026 = 令和8年度課税）"
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)
        f.write("\n")

    print("OK: 受入額 %s〜%s年度 / 控除額 %s〜%s課税。%d件の自治体の値が変わりました"
          % (labels[0], labels[-1], year_label(kojo_years[0]), year_label(kojo_years[-1]), changed))


if __name__ == "__main__":
    main()
