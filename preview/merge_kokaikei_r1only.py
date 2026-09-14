# -*- coding: utf-8 -*-
"""
R01年度ファイル（R1指標／H30指標の2シート）から、"R1指標"シートのみを
既存kokaikei.jsonに追記する専用スクリプト。
H30分は財政タブの開始年度（令和元年度）に合わせるため取り込まない。

使い方:
    python3 merge_kokaikei_r1only.py <既存kokaikei.json> <市町村指標一覧R01.xlsx> <都道府県指標一覧R01.xlsx> <出力先.json>
"""
import json
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

DATA_FILES = [
    "data-hokkaido-tohoku.json", "data-kanto.json", "data-chubu.json",
    "data-kinki.json", "data-chugoku-shikoku.json", "data-kyushu.json",
]

INDICATORS = [
    ("ka1", 7, 6), ("ka2", 10, 9), ("ka3", 13, 12), ("ka4", 16, 15),
    ("ka5", 19, 18), ("ka6", 22, 21), ("ka7", 25, 24), ("ka8", 28, 27),
    ("ka9", 31, 30),
]


def norm(name):
    if name is None:
        return ""
    return (str(name).replace("ヶ", "ケ").replace("ヵ", "カ")
                      .replace("\u3000", "").strip())


def read_muni_sheet(ws):
    out = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row or row[1] is None or row[2] is None:
            continue
        pref = norm(row[1]); name = norm(row[2])
        vals = {}
        for code, muni_col, _pc in INDICATORS:
            v = row[muni_col - 1]
            if isinstance(v, (int, float)):
                vals[code] = round(v, 1)
        out[(pref, name)] = vals
    return out


def read_pref_sheet(ws):
    out = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row or row[1] is None:
            continue
        pref = norm(row[1])
        vals = {}
        for code, _mc, pref_col in INDICATORS:
            v = row[pref_col - 1]
            if isinstance(v, (int, float)):
                vals[code] = round(v, 1)
        out[pref] = vals
    return out


def next_free_slot(entry, code):
    n = 1
    while (code + "_r" + str(n)) in entry:
        n += 1
    return n


def main():
    base_path, muni_path, pref_path, out_path = sys.argv[1:5]
    kokaikei = json.load(open(base_path, encoding="utf-8"))

    wb_m = openpyxl.load_workbook(muni_path, data_only=True)
    muni_r1 = read_muni_sheet(wb_m["R1指標"])
    wb_p = openpyxl.load_workbook(pref_path, data_only=True)
    pref_r1 = read_pref_sheet(wb_p["R1指標"])

    updated = 0
    for fn in DATA_FILES:
        db = json.load(open(os.path.join(ROOT, fn), encoding="utf-8"))
        for k, v in db.items():
            entry = kokaikei.get(k)
            if not entry:
                continue
            p = v.get("p")
            base_name = norm(re.sub(r"（[^）]*）$", "", k))
            is_pref = (p == k)
            nv = pref_r1.get(base_name) if is_pref else muni_r1.get((p, base_name))
            if not nv:
                continue
            for code, _mc, _pc in INDICATORS:
                slot = next_free_slot(entry, code)
                if nv.get(code) is not None:
                    entry[code + "_r" + str(slot)] = nv[code]
            updated += 1

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(kokaikei, f, ensure_ascii=False, separators=(",", ":"))
    print("更新: %d件 / %s に書き込みました" % (updated, out_path))


if __name__ == "__main__":
    main()
