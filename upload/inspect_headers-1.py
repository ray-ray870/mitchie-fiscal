# -*- coding: utf-8 -*-
"""
（確認専用・一時的なスクリプト）総務省のExcelの見出しを表示する。データは書き換えない。
人口ファイル（人口動態の列があるか）と、決算状況調（歳出の内訳の列があるか）を調べる。
"""
import json
import os
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
UA = {"User-Agent": "Mozilla/5.0 (mitchie-fiscal)"}


def out(t=""):
    print(t)
    sp = os.environ.get("GITHUB_STEP_SUMMARY")
    if sp:
        with open(sp, "a", encoding="utf-8") as fp:
            fp.write(t + "\n")


def show(label, url, head_rows, sample="姶良市", name_col=None):
    """見出しの行をまとめて、列ごとに「列番号：見出し ＝ 姶良市の値」を1行ずつ表示する"""
    out(f"## {label}\n- {url}")
    try:
        req = urllib.request.Request(url, headers=UA)
        data = urllib.request.urlopen(req, timeout=120).read()
    except Exception as ex:
        out(f"- 読めませんでした：{ex}\n")
        return
    p = Path("/tmp") / (label + ".xlsx")
    p.write_bytes(data)
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    out(f"- シート：{wb.sheetnames}")
    rows = list(ws.iter_rows(values_only=True, max_row=3000))
    width = max(len(r) for r in rows[:head_rows[1]])
    heads = []
    for j in range(width):
        parts = []
        for i in range(head_rows[0] - 1, head_rows[1]):
            v = rows[i][j] if j < len(rows[i]) else None
            if v not in (None, ""):
                parts.append(str(v).replace("\n", "").replace(" ", "").replace("\u3000", ""))
        heads.append("／".join(parts))
    sample_row = next((r for r in rows if any(isinstance(v, str) and v.replace("\u3000", "").strip() == sample for v in r)), None)
    out("```")
    for j in range(width):
        v = sample_row[j] if sample_row and j < len(sample_row) else None
        if heads[j] or v not in (None, ""):
            out(f"列{j}：{heads[j]} ＝ {v}")
    out("```\n")


def main():
    cfg = json.loads((ROOT / "scripts" / "config.json").read_text(encoding="utf-8"))
    show("人口（住民基本台帳）", cfg["population"]["muni"], (2, 6))
    show("目的別歳出（市）", cfg["purpose_expenditure"]["city"], (7, 13))
    show("性質別歳出？（市・番号から予想）", "https://www.soumu.go.jp/main_content/001061670.xlsx", (7, 13))
    show("性質別歳出？（町村・番号から予想）", "https://www.soumu.go.jp/main_content/001061675.xlsx", (7, 13), sample="湧水町")


if __name__ == "__main__":
    main()
