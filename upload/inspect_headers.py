# -*- coding: utf-8 -*-
"""
（確認専用・一時的なスクリプト）総務省のExcelの見出しを表示する。データは書き換えない。
人口ファイル（人口動態の列があるか）と、決算状況調（歳出の内訳の列があるか）を調べる。
"""
import json
import os
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
UA = {"User-Agent": "Mozilla/5.0 (mitchie-fiscal)"}


def out(t=""):
    print(t)
    sp = os.environ.get("GITHUB_STEP_SUMMARY")
    if sp:
        with open(sp, "a", encoding="utf-8") as fp:
            fp.write(t + "\n")


def show(label, url, max_row, sample_names=("姶良市", "夕張市", "鹿児島県", "北海道")):
    out(f"## {label}\n- {url}")
    req = urllib.request.Request(url, headers=UA)
    data = urllib.request.urlopen(req, timeout=120).read()
    p = Path("/tmp") / (label + ".xlsx")
    p.write_bytes(data)
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    out(f"- シート：{wb.sheetnames}")
    ws = wb[wb.sheetnames[0]]
    out("```")
    shown = 0
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = [f"{j}:{str(v).replace(chr(10), ' ')[:24]}" for j, v in enumerate(row) if v not in (None, "")]
        if i <= max_row:
            out(f"行{i} " + " | ".join(cells))
        elif any(isinstance(v, str) and any(s in v for s in sample_names) for v in row) and shown < 3:
            out(f"行{i} " + " | ".join(cells))
            shown += 1
        if i > 3000:
            break
    out("```\n")


def main():
    cfg = json.loads((ROOT / "scripts" / "config.json").read_text(encoding="utf-8"))
    show("人口（住民基本台帳）", cfg["population"]["muni"], 8)
    show("決算状況調（市）", cfg["budget"]["city"], 15)
    show("決算状況調（都道府県）", cfg["budget"]["pref"], 6)
    show("目的別歳出（市）", cfg["purpose_expenditure"]["city"], 15)


if __name__ == "__main__":
    main()
